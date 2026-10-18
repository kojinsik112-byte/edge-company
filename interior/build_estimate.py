# -*- coding: utf-8 -*-
"""엣지컴퍼니 인테리어 견적서 빌더 — estimate_data.py → 엑셀(수식) + HTML + PDF

사용: python -X utf8 interior/build_estimate.py
"""
import subprocess
from decimal import Decimal, ROUND_HALF_UP
from html import escape
from pathlib import Path
from urllib.parse import quote

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

import estimate_data as D

HERE = Path(__file__).resolve().parent
BASE = "엣지컴퍼니_인테리어견적서_32평올수리"
XLSX, HTML, PDF = (HERE / f"{BASE}.{e}" for e in ("xlsx", "html", "pdf"))
FONT_DIR = "../캡컷 ai편집기/assets/fonts"
CHROME = r"C:\Program Files\Google\Chrome\Application\chrome.exe"
NDAYS = 36


def xround(x, n=0):
    """엑셀 ROUND 와 같은 사사오입 (n<0 이면 십·백·천 단위)."""
    q = Decimal(1).scaleb(-n)
    return int(Decimal(str(x)).quantize(q, rounding=ROUND_HALF_UP))


def won_hangul(n):
    digits, small, big = "영일이삼사오육칠팔구", ["", "십", "백", "천"], ["", "만", "억", "조"]
    out, i = "", 0
    while n:
        g, n = n % 10000, n // 10000
        if g:
            out = "".join(digits[(g // 10 ** k) % 10] + small[k]
                          for k in range(3, -1, -1) if (g // 10 ** k) % 10) + big[i] + out
        i += 1
    return out


def compute():
    trades = []
    for code, name, summary, items in D.TRADES:
        rows = [(nm, spec, u, q, p, xround(q * p), note) for nm, spec, u, q, p, note in items]
        trades.append(dict(code=code, name=name, summary=summary, rows=rows,
                           sub=sum(r[5] for r in rows)))
    direct = sum(t["sub"] for t in trades)
    overhead = xround(direct * D.OVERHEAD_RATE, -3)
    supply = D.TARGET_TOTAL * 10 // 11
    adjust = supply - (direct + overhead)
    vat = xround(supply / 10)
    assert supply + vat == D.TARGET_TOTAL
    assert -500_000 < adjust <= 0, f"단수정리 {adjust:,} — 단가 조정 필요"
    return dict(trades=trades, direct=direct, overhead=overhead, subtotal=direct + overhead,
                adjust=adjust, supply=supply, vat=vat, total=supply + vat)


# ───────────────────────────── 엑셀 ─────────────────────────────
FN = "맑은 고딕"
INK, GOLD, BLUE = "1D1D1F", "A8875A", "0000FF"
thin = Side(style="thin", color="C9C2B8")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
F_HEAD = PatternFill("solid", fgColor="EFEAE2")
F_SEC = PatternFill("solid", fgColor="F7F4EF")
F_SUB = PatternFill("solid", fgColor="F2EEE7")
F_INPUT = PatternFill("solid", fgColor="FFF3B0")
F_TOTAL = PatternFill("solid", fgColor="E8DFD0")
MONEY = '#,##0;-#,##0;"-"'
MONEY_FREE = '#,##0;-#,##0;"무상"'
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")


def font(size=10, bold=False, color=INK):
    return Font(name=FN, size=size, bold=bold, color=color)


def put(ws, ref, value, *, f=None, fill=None, al=LEFT, fmt=None, border=True):
    c = ws[ref]
    c.value = value
    c.font = f or font()
    c.alignment = al
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    if border:
        c.border = BOX
    return c


def merge(ws, rng, value, **kw):
    first = rng.split(":")[0]
    put(ws, first, value, **kw)
    for row in ws[rng]:
        for c in row:
            if kw.get("border", True):
                c.border = BOX
            if kw.get("fill"):
                c.fill = kw["fill"]
    ws.merge_cells(rng)


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[L(i)].width = w


def a4(ws, landscape=False, title_rows=None):
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    if title_rows:
        ws.print_title_rows = title_rows
    ws.sheet_view.showGridLines = False


def sheet_detail(wb, C):
    ws = wb.create_sheet("상세내역")
    widths(ws, [7, 24, 46, 6, 8, 12, 14, 14])
    merge(ws, "A1:H1", "상 세  내 역 서", f=font(16, True), al=CENTER, border=False)
    ws.row_dimensions[1].height = 30
    merge(ws, "A2:H2", "산출 기준 — " + D.META["basis"], f=font(9, color="6E6E73"), border=False)
    ws.row_dimensions[2].height = 28
    merge(ws, "A3:H3", "단가 = 자재비 + 시공비 + 운반비 포함.  파란 글씨(수량·단가)를 고치면 금액·소계·견적서 합계가 자동으로 바뀝니다.",
          f=font(9, color=BLUE), border=False)
    for i, h in enumerate(["No", "품 명", "규격 · 사양", "단위", "수량", "단가(원)", "금액(원)", "비 고"]):
        put(ws, f"{L(i + 1)}4", h, f=font(10, True), fill=F_HEAD, al=CENTER)
    r, subs = 5, {}
    for t in C["trades"]:
        put(ws, f"A{r}", t["code"], f=font(10, True), fill=F_SEC, al=CENTER)
        merge(ws, f"B{r}:H{r}", t["name"], f=font(10, True), fill=F_SEC)
        r += 1
        first = r
        for i, (nm, spec, u, q, p, _amt, note) in enumerate(t["rows"], 1):
            put(ws, f"A{r}", f"{t['code']}-{i:02d}", f=font(9, color="6E6E73"), al=CENTER)
            put(ws, f"B{r}", nm)
            put(ws, f"C{r}", spec, f=font(9))
            put(ws, f"D{r}", u, al=CENTER)
            put(ws, f"E{r}", q, f=font(10, color=BLUE), al=RIGHT, fmt="General")
            put(ws, f"F{r}", p, f=font(10, color=BLUE), al=RIGHT, fmt=MONEY)
            put(ws, f"G{r}", f"=ROUND(E{r}*F{r},0)", al=RIGHT, fmt=MONEY_FREE)
            put(ws, f"H{r}", note, f=font(9))
            r += 1
        merge(ws, f"A{r}:F{r}", f"{t['name']}  소계", f=font(10, True), fill=F_SUB, al=CENTER)
        put(ws, f"G{r}", f"=SUM(G{first}:G{r - 1})", f=font(10, True), fill=F_SUB, al=RIGHT, fmt=MONEY)
        put(ws, f"H{r}", "", fill=F_SUB)
        subs[t["code"]] = r
        r += 1
    merge(ws, f"A{r}:F{r}", "직접공사비 합계", f=font(11, True), fill=F_TOTAL, al=CENTER)
    put(ws, f"G{r}", "=" + "+".join(f"G{x}" for x in subs.values()), f=font(11, True),
        fill=F_TOTAL, al=RIGHT, fmt=MONEY)
    put(ws, f"H{r}", "", fill=F_TOTAL)
    ws.freeze_panes = "A5"
    a4(ws, title_rows="4:4")
    return subs


def sheet_cover(wb, C, subs):
    ws = wb.active
    ws.title = "견적서"
    widths(ws, [6, 20, 44, 17, 10, 16])
    merge(ws, "A1:F1", "인 테 리 어   공 사   견 적 서", f=font(20, True), al=CENTER, border=False)
    ws.row_dimensions[1].height = 40
    merge(ws, "A2:F2", "EDGE COMPANY Co., Ltd.  ·  INTERIOR ESTIMATE", f=font(9, True, GOLD), al=CENTER, border=False)

    put(ws, "B4", "고 객 명", f=font(10, True), fill=F_HEAD, al=CENTER)
    merge(ws, "C4:D4", None, fill=F_INPUT, f=font(11, True, BLUE))
    ws["C4"].comment = Comment("고객 성함을 입력하세요.", "엣지컴퍼니")
    merge(ws, "E4:F4", "고객님 귀하", f=font(11, True))
    put(ws, "B5", "현장 주소", f=font(10, True), fill=F_HEAD, al=CENTER)
    merge(ws, "C5:F5", None, fill=F_INPUT, f=font(10, color=BLUE))
    ws["C5"].comment = Comment("현장 주소(아파트명·동·호수)를 입력하세요.", "엣지컴퍼니")
    info = [("공 사 명", D.META["title"]), ("공사 규모", D.META["scale"]), ("공사 기간", D.META["period"]),
            ("견적 번호", f"{D.META['no']}   (견적일 {D.META['date']})"), ("유효 기간", D.META["valid"])]
    for i, (k, v) in enumerate(info, 6):
        put(ws, f"B{i}", k, f=font(10, True), fill=F_HEAD, al=CENTER)
        merge(ws, f"C{i}:F{i}", v)

    merge(ws, "A12:F12", "공  급  자", f=font(10, True), fill=F_HEAD, al=CENTER)
    for i, (k, v) in enumerate(D.SUPPLIER, 13):
        put(ws, f"B{i}", k, f=font(10, True), fill=F_SEC, al=CENTER)
        merge(ws, f"C{i}:F{i}", v)

    top = 25  # 집계표 시작
    rows = {t["code"]: top + 1 + i for i, t in enumerate(C["trades"])}
    r_direct = top + 1 + len(C["trades"])
    r_oh, r_sub, r_adj, r_sup, r_vat, r_tot = (r_direct + k for k in range(1, 7))

    put(ws, "B20", "총 견적금액\n(부가세 포함)", f=font(11, True), fill=F_TOTAL, al=CENTER)
    ws.row_dimensions[20].height = 36
    put(ws, "C20", f"일금 {won_hangul(C['total'])}원정", f=font(15, True), fill=F_TOTAL, al=CENTER)
    ws["C20"].comment = Comment("총액이 바뀌면 한글 금액도 함께 고쳐 주세요.", "엣지컴퍼니")
    merge(ws, "D20:F20", f"=D{r_tot}", f=font(15, True), fill=F_TOTAL, al=CENTER, fmt='"₩"#,##0')
    merge(ws, "B21:F21", f'="공급가액 "&TEXT(D{r_sup},"#,##0")&"원  +  부가가치세 "&TEXT(D{r_vat},"#,##0")&"원"',
          f=font(9, color="6E6E73"), al=CENTER)
    merge(ws, "A23:F23", "아래와 같이 견적합니다.", f=font(10), border=False)

    for i, h in enumerate(["No", "공 종", "주요 내용", "금액(원)", "구성비", "비 고"]):
        put(ws, f"{L(i + 1)}{top}", h, f=font(10, True), fill=F_HEAD, al=CENTER)
    for t in C["trades"]:
        r = rows[t["code"]]
        put(ws, f"A{r}", t["code"], al=CENTER)
        put(ws, f"B{r}", t["name"], f=font(10, True))
        put(ws, f"C{r}", t["summary"], f=font(9))
        put(ws, f"D{r}", f"='상세내역'!G{subs[t['code']]}", f=font(10, color="008000"), al=RIGHT, fmt=MONEY)
        put(ws, f"E{r}", f"=IF($D${r_direct}=0,0,D{r}/$D${r_direct})", al=RIGHT, fmt="0.0%")
        put(ws, f"F{r}", "면허 직영" if t["code"] == "04" else "", f=font(9), al=CENTER)
    first, last = rows["01"], rows[C["trades"][-1]["code"]]

    def total_row(r, label, formula, *, note="", fill=F_SUB, f=None, fmt=MONEY):
        put(ws, f"A{r}", "", fill=fill)
        merge(ws, f"B{r}:C{r}", label, f=f or font(10, True), fill=fill)
        put(ws, f"D{r}", formula, f=f or font(10, True), fill=fill, al=RIGHT, fmt=fmt)
        put(ws, f"E{r}", "", fill=fill)
        put(ws, f"F{r}", note, f=font(9), fill=fill, al=CENTER)

    total_row(r_direct, "직접공사비 계", f"=SUM(D{first}:D{last})")
    ws[f"E{r_direct}"] = f"=SUM(E{first}:E{last})"
    ws[f"E{r_direct}"].number_format = "0.0%"
    total_row(r_oh, "공과잡비·현장관리비  (소모자재·양중·주차·현장관리)", f"=ROUND(D{r_direct}*E{r_oh},-3)",
              note="직접공사비 × 요율", fill=None)
    put(ws, f"E{r_oh}", D.OVERHEAD_RATE, f=font(10, color=BLUE), al=RIGHT, fmt="0.0%")
    ws[f"E{r_oh}"].comment = Comment("요율 입력값 (기본 5%).", "엣지컴퍼니")
    total_row(r_sub, "소   계", f"=D{r_direct}+D{r_oh}")
    total_row(r_adj, "단수 정리", C["adjust"], note="", fill=None)
    ws[f"D{r_adj}"].font = font(10, True, BLUE)
    ws[f"D{r_adj}"].comment = Comment("입력값. 수량·단가를 바꾸면 총액이 원하는 금액이 되도록 이 값을 조정하세요.", "엣지컴퍼니")
    total_row(r_sup, "공급가액", f"=D{r_sub}+D{r_adj}")
    total_row(r_vat, "부가가치세 (10%)", f"=ROUND(D{r_sup}/10,0)", fill=None)
    total_row(r_tot, "총 견적금액  (부가세 포함)", f"=D{r_sup}+D{r_vat}", fill=F_TOTAL, f=font(12, True))

    r = r_tot + 2
    for line in ["※ 파란 글씨 = 수정 가능한 입력값 (상세내역의 수량·단가, 요율, 단수 정리)  /  노란 칸 = 고객 정보 입력",
                 "※ 초록 글씨 = 상세내역 시트 소계를 끌어온 값 (직접 고치지 마세요)",
                 "※ 선택공사·공정표·결제 조건은 다음 시트에 있습니다."]:
        merge(ws, f"A{r}:F{r}", line, f=font(9, color="6E6E73"), border=False)
        r += 1
    a4(ws)
    return r_tot


def sheet_gantt(wb):
    ws = wb.create_sheet("공정표")
    widths(ws, [16, 40, 7] + [3.6] * NDAYS)
    lastcol = L(3 + NDAYS)
    merge(ws, f"A1:{lastcol}1", "공 정 표  —  6주 (작업일 36일 · 주 6일 작업, 일요일·공휴일 휴무)",
          f=font(15, True), al=CENTER, border=False)
    ws.row_dimensions[1].height = 30
    put(ws, "A2", "착공일 (월요일)", f=font(10, True), fill=F_HEAD, al=CENTER)
    y, m, d = D.META["start_example"]
    import datetime
    put(ws, "B2", datetime.date(y, m, d), f=font(10, True, BLUE), fill=F_INPUT, al=CENTER, fmt="yyyy-mm-dd")
    ws["B2"].comment = Comment("착공일(월요일)을 입력하면 날짜 행이 자동 계산됩니다. 공휴일은 수동으로 순연하세요.", "엣지컴퍼니")
    merge(ws, f"C2:{lastcol}2", "  ← 착공일을 입력하면 아래 날짜가 바뀝니다 (예시 날짜)", f=font(9, color="6E6E73"), border=False)

    for c, h in (("A", "공 정"), ("B", "세부 내용"), ("C", "일수")):
        merge(ws, f"{c}4:{c}6", h, f=font(10, True), fill=F_HEAD, al=CENTER)
    weekdays = "월화수목금토"
    for w in range(6):
        c0, c1 = L(4 + w * 6), L(9 + w * 6)
        merge(ws, f"{c0}4:{c1}4", f"{w + 1}주차", f=font(9, True), fill=F_HEAD, al=CENTER)
    for k in range(1, NDAYS + 1):
        col = L(3 + k)
        put(ws, f"{col}5", f"D{k}", f=font(7, True), fill=F_HEAD, al=CENTER)
        put(ws, f"{col}6", weekdays[(k - 1) % 6], f=font(7), fill=F_HEAD, al=CENTER)
        put(ws, f"{col}7", f"=$B$2+{k - 1}+INT({k - 1}/6)", f=font(6, color="6E6E73"), al=CENTER, fmt="m/d")
    put(ws, "A7", "날짜", f=font(8, color="6E6E73"), al=CENTER)
    put(ws, "B7", "", border=True)
    put(ws, "C7", "", border=True)

    r = 8
    for name, detail, s, e, cat in D.GANTT:
        put(ws, f"A{r}", name, f=font(10, True))
        put(ws, f"B{r}", detail, f=font(9))
        put(ws, f"C{r}", e - s + 1, al=CENTER)
        color = D.GANTT_CATS[cat][1]
        for k in range(1, NDAYS + 1):
            c = put(ws, f"{L(3 + k)}{r}", None, al=CENTER)
            if s <= k <= e:
                c.fill = PatternFill("solid", fgColor=color)
        ws.row_dimensions[r].height = 18
        r += 1
    put(ws, f"A{r}", "★ 고객 확인", f=font(10, True, "B23B3B"))
    put(ws, f"B{r}", "현장 확인·승인 시점 (아래 표 참고)", f=font(9))
    put(ws, f"C{r}", "")
    for k in range(1, NDAYS + 1):
        put(ws, f"{L(3 + k)}{r}", "★" if k in D.CHECK_DAYS else None, f=font(9, True, "B23B3B"), al=CENTER)
    r += 1
    put(ws, f"A{r}", "₩ 결제", f=font(10, True))
    put(ws, f"B{r}", "D1 착공금 30% · D12 중도금 40% · D36 잔금 20%", f=font(9))
    put(ws, f"C{r}", "")
    for k in range(1, NDAYS + 1):
        put(ws, f"{L(3 + k)}{r}", "₩" if k in D.PAY_DAYS else None, f=font(9, True), al=CENTER)
    r += 2
    put(ws, f"A{r}", "범 례", f=font(10, True), fill=F_HEAD, al=CENTER)
    col = 2
    for label, color in D.GANTT_CATS.values():
        pass
    legend = "   ".join(f"■ {label}" for label, _ in D.GANTT_CATS.values())
    merge(ws, f"B{r}:{lastcol}{r}", legend, f=font(9))
    r += 2
    put(ws, f"A{r}", "시 점", f=font(10, True), fill=F_HEAD, al=CENTER)
    merge(ws, f"B{r}:{lastcol}{r}", "고객 확인 포인트", f=font(10, True), fill=F_HEAD, al=CENTER)
    for when, text in D.CHECKPOINTS:
        r += 1
        put(ws, f"A{r}", when, f=font(10, True), al=CENTER)
        merge(ws, f"B{r}:{lastcol}{r}", text, f=font(9))
    a4(ws, landscape=True)


def sheet_options(wb):
    ws = wb.create_sheet("선택공사")
    widths(ws, [5, 22, 44, 26, 30])
    merge(ws, "A1:E1", "선 택 공 사  (견적 금액 미포함 · 참고 예상가)", f=font(15, True), al=CENTER, border=False)
    ws.row_dimensions[1].height = 30
    merge(ws, "A2:E2", "착공 전에 결정하시면 배선·보강·일정에 미리 반영합니다. 금액은 실측·제품 선택 후 확정됩니다.",
          f=font(9, color="6E6E73"), border=False)
    for i, h in enumerate(["No", "항 목", "내 용", "참고 예상가", "비 고"]):
        put(ws, f"{L(i + 1)}4", h, f=font(10, True), fill=F_HEAD, al=CENTER)
    for i, (item, desc, price, note) in enumerate(D.OPTIONS, 1):
        r = 4 + i
        put(ws, f"A{r}", i, al=CENTER)
        put(ws, f"B{r}", item, f=font(10, True))
        put(ws, f"C{r}", desc, f=font(9))
        put(ws, f"D{r}", price, al=CENTER)
        put(ws, f"E{r}", note, f=font(9))
        ws.row_dimensions[r].height = 20
    a4(ws)


def sheet_terms(wb, r_tot):
    ws = wb.create_sheet("계약조건")
    widths(ws, [16, 36, 12, 18, 18])
    merge(ws, "A1:E1", "결 제 조 건  ·  특 기 사 항", f=font(15, True), al=CENTER, border=False)
    ws.row_dimensions[1].height = 30
    for i, h in enumerate(["단 계", "시 점", "비 율", "금액(원)"]):
        put(ws, f"{L(i + 1)}3", h, f=font(10, True), fill=F_HEAD, al=CENTER)
    r = 4
    for stage, when, ratio in D.PAYMENTS:
        put(ws, f"A{r}", stage, f=font(10, True), al=CENTER)
        put(ws, f"B{r}", when)
        put(ws, f"C{r}", ratio, f=font(10, color=BLUE), al=CENTER, fmt="0%")
        put(ws, f"D{r}", f"=ROUND('견적서'!$D${r_tot}*C{r},0)", al=RIGHT, fmt=MONEY)
        r += 1
    put(ws, f"A{r}", "합 계", f=font(10, True), fill=F_SUB, al=CENTER)
    put(ws, f"B{r}", "", fill=F_SUB)
    put(ws, f"C{r}", f"=SUM(C4:C{r - 1})", f=font(10, True), fill=F_SUB, al=CENTER, fmt="0%")
    put(ws, f"D{r}", f"=SUM(D4:D{r - 1})", f=font(10, True), fill=F_SUB, al=RIGHT, fmt=MONEY)
    r += 2

    def section(title, lines, numbered=False):
        nonlocal r
        merge(ws, f"A{r}:E{r}", title, f=font(11, True), fill=F_HEAD)
        r += 1
        for i, line in enumerate(lines, 1):
            merge(ws, f"A{r}:E{r}", (f"{i}. " if numbered else "· ") + line, f=font(9))
            ws.row_dimensions[r].height = 30 if len(line) > 60 else 18
            r += 1
        r += 1

    section("포함 사항", D.INCLUDED)
    section("불포함 사항", D.EXCLUDED)
    section("특기 사항", D.NOTES, numbered=True)
    merge(ws, f"A{r}:E{r}", "하자보수 (준공일로부터 무상 A/S)", f=font(11, True), fill=F_HEAD)
    r += 1
    for k, v in D.WARRANTY:
        put(ws, f"A{r}", "", border=False)
        merge(ws, f"A{r}:B{r}", k)
        put(ws, f"C{r}", v, f=font(10, True), al=CENTER)
        r += 1
    r += 1
    merge(ws, f"A{r}:E{r}", "위 견적 내용을 확인하였으며, 이에 따라 공사를 의뢰합니다.          20      년      월      일",
          f=font(10), border=False)
    r += 2
    merge(ws, f"A{r}:B{r}", "고객 (갑)  :                                   (서명)", f=font(10, True), border=False)
    merge(ws, f"C{r}:E{r}", "공급자 (을)  :  ㈜엣지컴퍼니 대표이사 고진식  (인)", f=font(10, True), border=False)
    a4(ws)


def build_xlsx(C):
    wb = Workbook()  # 첫 시트(active)는 sheet_cover 가 '견적서'로 채운다
    subs = sheet_detail(wb, C)
    r_tot = sheet_cover(wb, C, subs)
    sheet_gantt(wb)
    sheet_options(wb)
    sheet_terms(wb, r_tot)
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True
    wb.save(XLSX)


# ───────────────────────────── HTML / PDF ─────────────────────────────
def won(n):
    return f"{n:,}"


def build_html(C):
    e = escape
    M = D.META
    font_r = quote(f"{FONT_DIR}/Pretendard-Regular.otf")
    font_b = quote(f"{FONT_DIR}/Pretendard-Bold.otf")

    sum_rows = "".join(
        f"<tr><td class='c'>{t['code']}</td><td class='b'>{e(t['name'])}</td><td class='s'>{e(t['summary'])}</td>"
        f"<td class='r'>{won(t['sub'])}</td><td class='r'>{t['sub'] / C['direct'] * 100:.1f}%</td></tr>"
        for t in C["trades"])
    tot = [("직접공사비 계", C["direct"], "sub"),
           (f"공과잡비·현장관리비 ({D.OVERHEAD_RATE * 100:.0f}%)  <small>소모자재·양중·주차·현장관리</small>", C["overhead"], ""),
           ("단수 정리", C["adjust"], ""),
           ("공급가액", C["supply"], "sub"),
           ("부가가치세 (10%)", C["vat"], ""),
           ("총 견적금액 <small>(부가세 포함)</small>", C["total"], "grand")]
    sum_rows += "".join(f"<tr class='{cls}'><td></td><td colspan='2'>{lab}</td><td class='r'>{won(v)}</td><td></td></tr>"
                        for lab, v, cls in tot)

    detail = []
    for t in C["trades"]:
        detail.append(f"<tr class='sec'><td class='c'>{t['code']}</td><td colspan='7'>{e(t['name'])}</td></tr>")
        for i, (nm, spec, u, q, p, amt, note) in enumerate(t["rows"], 1):
            qs = f"{q:g}"
            detail.append(
                f"<tr><td class='c m'>{i}</td><td>{e(nm)}</td><td class='s'>{e(spec)}</td><td class='c'>{u}</td>"
                f"<td class='r'>{qs}</td><td class='r'>{won(p) if p else '-'}</td>"
                f"<td class='r'>{won(amt) if amt else '<b class=free>무상</b>'}</td><td class='s'>{e(note)}</td></tr>")
        detail.append(f"<tr class='sub'><td colspan='6' class='c'>{e(t['name'])} 소계</td>"
                      f"<td class='r'>{won(t['sub'])}</td><td></td></tr>")
    detail.append(f"<tr class='grand'><td colspan='6' class='c'>직접공사비 합계</td><td class='r'>{won(C['direct'])}</td><td></td></tr>")

    opts = "".join(f"<tr><td class='c'>{i}</td><td class='b'>{e(a)}</td><td class='s'>{e(b)}</td>"
                   f"<td class='c'>{e(c)}</td><td class='s'>{e(d)}</td></tr>"
                   for i, (a, b, c, d) in enumerate(D.OPTIONS, 1))
    pays = "".join(f"<tr><td class='c b'>{s}</td><td>{w}</td><td class='c'>{r * 100:.0f}%</td>"
                   f"<td class='r'>{won(xround(C['total'] * r))}</td></tr>" for s, w, r in D.PAYMENTS)

    # 공정표
    head_w = "".join(f"<th colspan='6' class='wk'>{w + 1}주차</th>" for w in range(6))
    head_d = "".join(f"<th class='d{' ws' if k % 6 == 1 else ''}'>D{k}</th>" for k in range(1, NDAYS + 1))
    head_y = "".join(f"<th class='d{' ws' if k % 6 == 1 else ''}'>{'월화수목금토'[(k - 1) % 6]}</th>" for k in range(1, NDAYS + 1))
    g_rows = []
    for name, det, s, en, cat in D.GANTT:
        color = D.GANTT_CATS[cat][1]
        cells = "".join(
            f"<td class='g{' ws' if k % 6 == 1 else ''}'"
            + (f" style='background:#{color}'" if s <= k <= en else "") + "></td>"
            for k in range(1, NDAYS + 1))
        g_rows.append(f"<tr><td class='b'>{e(name)}</td><td class='s'>{e(det)}</td><td class='c'>{en - s + 1}</td>{cells}</tr>")
    star = "".join(f"<td class='g mk{' ws' if k % 6 == 1 else ''}'>{'★' if k in D.CHECK_DAYS else ''}</td>" for k in range(1, NDAYS + 1))
    pay = "".join(f"<td class='g pk{' ws' if k % 6 == 1 else ''}'>{'₩' if k in D.PAY_DAYS else ''}</td>" for k in range(1, NDAYS + 1))
    g_rows.append(f"<tr class='mark'><td class='b red'>★ 고객 확인</td><td class='s'>현장 확인·승인 시점</td><td></td>{star}</tr>")
    g_rows.append(f"<tr class='mark'><td class='b'>₩ 결제</td><td class='s'>D1 착공금 30% · D12 중도금 40% · D36 잔금 20%</td><td></td>{pay}</tr>")
    legend = "".join(f"<span><i style='background:#{c}'></i>{e(l)}</span>" for l, c in D.GANTT_CATS.values())
    checks = "".join(f"<div class='ck'><b>{e(w)}</b><span>{e(t)}</span></div>" for w, t in D.CHECKPOINTS)

    inc = "".join(f"<li>{e(x)}</li>" for x in D.INCLUDED)
    exc = "".join(f"<li>{e(x)}</li>" for x in D.EXCLUDED)
    notes = "".join(f"<li>{e(x)}</li>" for x in D.NOTES)
    warr = "".join(f"<tr><td>{e(k)}</td><td class='c b'>{v}</td></tr>" for k, v in D.WARRANTY)
    supplier = "".join(f"<tr><th>{e(k)}</th><td>{e(v)}</td></tr>" for k, v in D.SUPPLIER)

    html = f"""<!doctype html><html lang="ko"><head><meta charset="utf-8">
<title>인테리어 공사 견적서 — {e(M['title'])}</title>
<style>
@font-face{{font-family:'PretendardE';src:url('{font_r}');font-weight:400}}
@font-face{{font-family:'PretendardE';src:url('{font_b}');font-weight:700}}
:root{{--ink:#1d1d1f;--muted:#6e6e73;--gold:#a8875a;--line:#d8d2c8;--head:#f1ece4;--sec:#f8f5f0;--sub:#f3efe8;--red:#b23b3b}}
@page{{size:A4;margin:13mm 12mm 15mm;
  @bottom-left{{content:"㈜엣지컴퍼니 · 인테리어 공사 견적서 {M['no']}";font:7.5pt 'PretendardE','Malgun Gothic';color:#9a958d}}
  @bottom-right{{content:counter(page) " / " counter(pages);font:7.5pt 'PretendardE','Malgun Gothic';color:#9a958d}}}}
@page land{{size:A4 landscape;margin:11mm 11mm 13mm;
  @bottom-left{{content:"㈜엣지컴퍼니 · 인테리어 공사 견적서 {M['no']}";font:7.5pt 'PretendardE','Malgun Gothic';color:#9a958d}}
  @bottom-right{{content:counter(page) " / " counter(pages);font:7.5pt 'PretendardE','Malgun Gothic';color:#9a958d}}}}
*{{box-sizing:border-box;-webkit-print-color-adjust:exact;print-color-adjust:exact}}
html,body{{margin:0;background:#fff;color:var(--ink);font-family:'PretendardE','Malgun Gothic',sans-serif;font-size:9pt;line-height:1.45}}
.page{{break-after:page}} .page:last-child{{break-after:auto}} .land{{page:land}}
h1{{font-size:24pt;letter-spacing:.32em;margin:2mm 0 1mm;font-weight:700;text-align:center}}
h2{{font-size:12.5pt;margin:0 0 2.5mm;padding-left:3mm;border-left:3px solid var(--gold);font-weight:700}}
h2 small,h1 small{{font-size:8.5pt;color:var(--muted);font-weight:400;letter-spacing:0;margin-left:2mm}}
.brand{{text-align:center;color:var(--gold);font-weight:700;font-size:8pt;letter-spacing:.28em}}
.docno{{text-align:center;color:var(--muted);font-size:8.5pt;margin-bottom:4mm}}
.client{{font-size:14pt;font-weight:700;margin:0 0 4mm}} .client .blank{{display:inline-block;width:62mm;border-bottom:1px solid var(--ink);margin-right:2mm}}
table{{width:100%;border-collapse:collapse}}
th,td{{border:1px solid var(--line);padding:1.3mm 2mm;vertical-align:middle}}
th{{background:var(--head);font-weight:700;text-align:center}}
.grid2{{display:grid;grid-template-columns:1.05fr 1fr;gap:4mm;margin-bottom:3.5mm}}
.kv th{{width:25mm;font-size:8pt;white-space:nowrap;padding:1mm 1.5mm}} .kv td{{font-size:8.2pt;padding:1mm 2mm}} .kv .fill{{height:7mm}}
.kv caption{{text-align:left;font-weight:700;font-size:9pt;padding-bottom:1.2mm;color:var(--gold)}}
.total{{display:grid;grid-template-columns:34mm 1fr auto;align-items:center;border:1.5px solid var(--ink);padding:3mm 5mm;margin:0 0 1.2mm}}
.total .lbl{{font-weight:700;font-size:10pt;line-height:1.3}} .total .lbl small{{display:block;font-weight:400;color:var(--muted);font-size:7.5pt}}
.total .wonk{{font-size:15pt;font-weight:700;letter-spacing:.08em}} .total .num{{font-size:19pt;font-weight:700}}
.split{{text-align:right;color:var(--muted);font-size:8pt;margin-bottom:2.5mm}}
.lead{{margin:0 0 2mm;font-size:9pt}}
td.c{{text-align:center}} td.r{{text-align:right;font-variant-numeric:tabular-nums;white-space:nowrap}} .b{{font-weight:700}} td.s{{font-size:8pt;color:#3a3a3c}} .m{{color:var(--muted)}}
.sum td{{font-size:8.5pt;padding:.85mm 2mm}} .sum tr.sub td{{background:var(--sub);font-weight:700}}
.sum small{{color:var(--muted);font-weight:400;font-size:7.5pt}}
tr.grand td{{background:#e8dfd0;font-weight:700;font-size:10pt}}
.detail td{{font-size:8.3pt;padding:1.05mm 1.6mm}} .detail thead th{{font-size:8.3pt}}
.detail tr{{break-inside:avoid}} .detail tr.sec td{{background:var(--sec);font-weight:700;border-top:1.5px solid var(--gold)}}
.detail tr.sub td{{background:var(--sub);font-weight:700}} .free{{color:var(--gold)}}
.basis{{font-size:8pt;color:var(--muted);margin:0 0 2.5mm}}
.note{{font-size:8pt;color:var(--muted);margin:1.5mm 0 5mm}}
.opt td{{font-size:8.4pt}} .pay td{{font-size:9pt}}
.two{{display:grid;grid-template-columns:1fr 1fr;gap:5mm}}
ul,ol{{margin:0;padding-left:5mm}} li{{margin:0 0 1.2mm;font-size:8.6pt}}
.gantt th,.gantt td{{padding:.6mm .8mm;font-size:7.4pt}} .gantt td.g{{padding:0;height:5.1mm}}
.gantt th.d{{font-size:6.3pt;padding:.4mm 0;font-weight:400}} .gantt th.wk{{font-size:7.5pt}}
.gantt .ws{{border-left:1.5px solid #9a958d}} .gantt tr.mark td{{background:#fbfaf7}}
.gantt td.mk{{color:var(--red);text-align:center;font-size:8pt}} .gantt td.pk{{text-align:center;font-size:7.5pt;font-weight:700}} .red{{color:var(--red)}}
.legend{{display:flex;gap:5mm;flex-wrap:wrap;margin:2mm 0 3mm;font-size:8pt}} .legend i{{display:inline-block;width:4mm;height:3mm;margin-right:1.2mm;vertical-align:-.3mm}}
.cks{{display:grid;grid-template-columns:1fr 1fr;gap:1.2mm 5mm}} .ck{{display:flex;gap:2mm;font-size:8pt;border-bottom:1px dotted var(--line);padding-bottom:1mm}} .ck b{{min-width:12mm;color:var(--red)}}
.warr td{{font-size:9pt}}
.sign{{margin-top:7mm;border-top:1.5px solid var(--ink);padding-top:4mm}} .sign p{{margin:0 0 6mm}}
.sign .row{{display:grid;grid-template-columns:1fr 1fr;gap:8mm;font-size:10pt}} .sign .ln{{display:inline-block;width:45mm;border-bottom:1px solid var(--ink)}}
.sign .date{{text-align:right;margin-top:6mm;letter-spacing:.1em}}
</style></head><body>

<section class="page">
  <div class="brand">EDGE COMPANY · INTERIOR ESTIMATE</div>
  <h1>인테리어 공사 견적서</h1>
  <div class="docno">견적번호 {M['no']} &nbsp;·&nbsp; 견적일 {M['date']}</div>
  <div class="client"><span class="blank"></span>고객님 귀하</div>
  <div class="grid2">
    <table class="kv"><caption>공사 개요</caption>
      <tr><th>공 사 명</th><td class="b">{e(M['title'])}</td></tr>
      <tr><th>현장 주소</th><td class="fill"></td></tr>
      <tr><th>공사 규모</th><td>{e(M['scale'])}</td></tr>
      <tr><th>공사 기간</th><td>{e(M['period'])}</td></tr>
      <tr><th>유효 기간</th><td>{e(M['valid'])}</td></tr>
    </table>
    <table class="kv"><caption>공 급 자</caption>{supplier}</table>
  </div>
  <div class="total"><div class="lbl">총 견적금액<small>부가가치세 포함</small></div>
    <div class="wonk">일금 {won_hangul(C['total'])}원정</div><div class="num">₩{won(C['total'])}</div></div>
  <div class="split">공급가액 {won(C['supply'])}원 + 부가가치세 {won(C['vat'])}원</div>
  <p class="lead">아래와 같이 견적합니다.</p>
  <table class="sum"><thead><tr><th style="width:9mm">No</th><th style="width:38mm">공 종</th><th>주요 내용</th><th style="width:27mm">금액(원)</th><th style="width:16mm">구성비</th></tr></thead>
  <tbody>{sum_rows}</tbody></table>
</section>

<section class="page">
  <h2>상세 내역서 <small>단가 = 자재비 + 시공비 + 운반비 포함</small></h2>
  <p class="basis">산출 기준 — {e(M['basis'])}</p>
  <table class="detail"><thead><tr><th style="width:8mm">No</th><th style="width:36mm">품 명</th><th>규격 · 사양</th><th style="width:10mm">단위</th>
  <th style="width:11mm">수량</th><th style="width:19mm">단가(원)</th><th style="width:21mm">금액(원)</th><th style="width:21mm">비 고</th></tr></thead>
  <tbody>{''.join(detail)}</tbody></table>
</section>

<section class="page">
  <h2>선택공사 <small>견적 금액 미포함 · 참고 예상가</small></h2>
  <table class="opt"><thead><tr><th style="width:8mm">No</th><th style="width:36mm">항 목</th><th>내 용</th><th style="width:44mm">참고 예상가</th><th style="width:46mm">비 고</th></tr></thead><tbody>{opts}</tbody></table>
  <p class="note">착공 전에 결정하시면 배선·천장 보강·일정에 미리 반영합니다. 금액은 실측·제품 선택 후 확정됩니다.</p>
  <h2>결제 조건</h2>
  <table class="pay"><thead><tr><th style="width:24mm">단 계</th><th>시 점</th><th style="width:20mm">비 율</th><th style="width:36mm">금액(원)</th></tr></thead>
  <tbody>{pays}<tr class="grand"><td class="c">합 계</td><td></td><td class="c">100%</td><td class="r">{won(C['total'])}</td></tr></tbody></table>
  <p class="note">잔금은 고객 동행 점검 후 지적 사항 보수가 끝난 뒤 받습니다.</p>
  <div class="two">
    <div><h2>포함 사항</h2><ul>{inc}</ul></div>
    <div><h2>불포함 사항</h2><ul>{exc}</ul></div>
  </div>
</section>

<section class="page land">
  <h2>공정표 <small>{e(M['period'])} · 주 6일(월~토) 작업, 착공일 확정 시 날짜 기입</small></h2>
  <table class="gantt"><thead>
    <tr><th rowspan="2" style="width:25mm">공 정</th><th rowspan="2" style="width:58mm">세부 내용</th><th rowspan="2" style="width:8mm">일수</th>{head_w}</tr>
    <tr>{head_d}</tr><tr><th></th><th></th><th></th>{head_y}</tr></thead>
  <tbody>{''.join(g_rows)}</tbody></table>
  <div class="legend">{legend}</div>
  <div class="cks">{checks}</div>
</section>

<section class="page">
  <h2>특기 사항</h2>
  <ol>{notes}</ol>
  <div style="height:5mm"></div>
  <div class="two">
    <div><h2>하자보수 <small>준공일로부터 무상 A/S</small></h2>
      <table class="warr"><thead><tr><th>공 종</th><th style="width:22mm">기 간</th></tr></thead><tbody>{warr}</tbody></table></div>
    <div><h2>시공 관리</h2><ul>
      <li>공정별 시공 사진·영상 보고 (주 2회 이상)</li>
      <li>철거 후 숨은 하자는 사진 공유 → 서면 합의 후 진행</li>
      <li>방수 후 48시간 담수 테스트 결과 보고</li>
      <li>준공 전 전기 안전점검(절연저항 측정) 결과 제공</li>
      <li>준공 시 A/S 보증서 발급</li></ul></div>
  </div>
  <div class="sign">
    <p>위 견적 내용을 확인하였으며, 이에 따라 공사를 의뢰합니다.</p>
    <div class="row"><div>고객 (갑) &nbsp; <span class="ln"></span> (서명)</div>
      <div>공급자 (을) &nbsp; ㈜엣지컴퍼니 대표이사 고진식 &nbsp;(인)</div></div>
    <div class="date">20 &nbsp;&nbsp;&nbsp;&nbsp; 년 &nbsp;&nbsp;&nbsp;&nbsp; 월 &nbsp;&nbsp;&nbsp;&nbsp; 일</div>
  </div>
</section>
</body></html>"""
    HTML.write_text(html, encoding="utf-8")


def build_pdf():
    subprocess.run([CHROME, "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
                    "--virtual-time-budget=5000", f"--print-to-pdf={PDF}", HTML.as_uri()],
                   check=True, capture_output=True, timeout=120)


if __name__ == "__main__":
    C = compute()
    build_xlsx(C)
    build_html(C)
    build_pdf()
    print(f"직접공사비 {C['direct']:,} / 관리비 {C['overhead']:,} / 단수정리 {C['adjust']:,}")
    print(f"공급가액 {C['supply']:,} + 부가세 {C['vat']:,} = 총액 {C['total']:,}")
    print(XLSX.name, HTML.name, PDF.name, "생성 완료")
