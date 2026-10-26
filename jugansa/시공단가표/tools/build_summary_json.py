"""엑셀(업체별_작성용, 재계산본) → 품목별 범위(최저~최고) JSON. build_handout.py 입력용."""
import json, sys, re
import openpyxl

XLSX, DATA, OUT = sys.argv[1], sys.argv[2], sys.argv[3]
data = json.load(open(DATA))
ws = openpyxl.load_workbook(XLSX, data_only=True)["업체별_작성용"]
rows = {}
for r in range(5, ws.max_row + 1):
    v = [ws.cell(r, c).value for c in range(1, 9)]
    if v[0]:
        rows.setdefault(v[0], []).append(v)


def rng(vals):
    nums, texts = [], []
    for x in vals:
        if x in (None, "", "-"):
            continue
        if isinstance(x, (int, float)):
            nums.append(x)
            continue
        mm = re.fullmatch(r"(\D*?)\s*([\d.]+)\s*만원", str(x).strip())
        if mm:
            texts.append((mm.group(1).strip(), float(mm.group(2))))
        else:
            texts.append((str(x), None))
    if nums:
        lo, hi = min(nums), max(nums)
        return f"{lo:,.0f}" if lo == hi else f"{lo:,.0f}~{hi:,.0f}"
    pref = [t for t in texts if t[1] is not None]
    if pref:
        lo, hi = min(t[1] for t in pref), max(t[1] for t in pref)
        f = lambda n: f"{n:g}"
        return f"{pref[0][0]} {f(lo)}만원" if lo == hi else f"{pref[0][0]} {f(lo)}~{f(hi)}만원"
    if texts:
        return texts[0][0]
    return "-" if any(x == "-" for x in vals) else None


first = lambda xs: next((x for x in xs if x not in (None, "", "-")), "")
groups = []
for g in data["groups"]:
    items = []
    for it in g["items"]:
        vs = rows.get(it["item"])
        if not vs:
            continue
        items.append(dict(item=it["item"], scope=first(v[3] for v in vs), normal=rng([v[4] for v in vs]),
                          normal_desc=first(v[5] for v in vs), premium=rng([v[6] for v in vs]),
                          premium_desc=first(v[7] for v in vs), option=f"참여 {len(vs)}개 업체" if len(vs) > 1 else ""))
    if items:
        groups.append(dict(name=g["name"], items=items))
json.dump(dict(groups=groups), open(OUT, "w"), ensure_ascii=False, indent=1)
for g in groups:
    for i in g["items"]:
        print(i["item"], "|", i["normal"], i["normal_desc"], "|", i["premium"], i["premium_desc"])
