"""입주박람회 시공 단가표 엑셀: 품목별 대표가(일반/프리미엄) + 업체별 A·B·C(대표가 자동연결) + 가전·가구 특가."""
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L

DB, DATA, OUT = sys.argv[1], sys.argv[2], sys.argv[3]
OVR = json.load(open(sys.argv[4]))["overrides"] if len(sys.argv) > 4 else {}
rows = json.load(open(DB))["rows"]
data = json.load(open(DATA))
EVENT = "E5"

FONT = "맑은 고딕"
NAVY = "12345E"
f_base = Font(name=FONT, size=10)
f_bold = Font(name=FONT, size=10, bold=True)
f_price = Font(name=FONT, size=11, bold=True, color="0000FF")
f_head = Font(name=FONT, size=10, bold=True, color="FFFFFF")
f_title = Font(name=FONT, size=16, bold=True, color=NAVY)
f_grp = Font(name=FONT, size=10, bold=True, color=NAVY)
f_item = Font(name=FONT, size=11, bold=True)
f_input = Font(name=FONT, size=10, color="0000FF")
f_link = Font(name=FONT, size=10, color="008000")
f_linkb = Font(name=FONT, size=11, bold=True, color="008000")
f_note = Font(name=FONT, size=9, color="666666", italic=True)
fill_head = PatternFill("solid", fgColor=NAVY)
fill_grp = PatternFill("solid", fgColor="E8F1FB")
fill_yel = PatternFill("solid", fgColor="FFFF00")
fill_ex = PatternFill("solid", fgColor="FFF4D6")
thin = Side(style="thin", color="B9C4D2")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left = Alignment(horizontal="left", vertical="center", wrap_text=True)
PRICE = '#,##0"만원";;"-";@'
PCT = '0%;;"-"'

wb = Workbook()


def cs(ws, r, c, v, font=f_base, align=center, fill=None, fmt=None):
    x = ws.cell(r, c, v)
    x.font, x.alignment, x.border = font, align, box
    if fill:
        x.fill = fill
    if fmt:
        x.number_format = fmt
    return x


def header(ws, r, labels):
    for i, t in enumerate(labels, 1):
        cs(ws, r, i, t, f_head, center, fill_head)


def widths(ws, w):
    for i, x in enumerate(w, 1):
        ws.column_dimensions[L(i)].width = x


def a4(ws, landscape=False):
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.4


# ======================= 안내 =======================
g = wb.active
g.title = "안내"
g["A1"] = "입주박람회 시공 단가표 — 사용 안내"
g["A1"].font = f_title
guide = [
    ("무엇인가", "입주민이 '시공 단가표 주세요' 할 때 보여주는 품목별 시공 범위·가격표. 업체들 가격이 비슷하므로 품목별 대표가 1줄로 정리."),
    ("고칠 곳", "'시공단가표' 시트의 파란 글씨(범위·가격)만 고치면 업체별 시트까지 자동 반영."),
    ("업체별 A·B·C", "'업체별_작성용'에 업체명이 있고, 범위·가격은 대표가가 자동으로 들어감(초록). 가격이 다른 업체만 그 칸에 숫자를 덮어쓰기. '업체별_공개용'은 업체명 대신 줄눈A·줄눈B…로 자동 표시."),
    ("가격 근거", "줄눈·탄성코트 = 본부장님 제시가. 나머지 = 공개 시세 범위 안에서 AI가 잡은 대표가(초안). '시세근거' 시트에 범위·출처."),
    ("배포 전 확인", "대표가는 업체가 확인한 금액이 아님 → 배포 전 품목별 업체에 '이 금액 맞나요' 1회 확인 필수(현장에서 가격이 다르면 입주민 민원)."),
    ("표시 기준", "84㎡(34평) 기준 · 현장 계약 시 금액 · 평형·현장 여건에 따라 달라질 수 있음. 부가세 포함 여부는 업체별로 통일해 표기."),
    ("가전·가구", "업체별 대표 특가 1~2개만 사전 공개, 그 외는 박람회 현장 공개. '가전가구_특가' 시트."),
    ("색상", "파란 글씨 = 입력값 / 초록 글씨 = 다른 시트에서 자동 연결 / 노란 칸 = 아직 비어 있는 입력칸"),
]
for i, (k, v) in enumerate(guide):
    cs(g, 3 + i, 1, k, f_bold, center, fill_grp)
    cs(g, 3 + i, 2, v, f_base, left)
    g.row_dimensions[3 + i].height = 40
widths(g, [14, 92])

# ======================= 시공단가표 (마스터) =======================
m = wb.create_sheet("시공단가표")
m["A1"] = "입주박람회 시공 단가표"
m["A1"].font = f_title
m["A2"] = "84㎡(34평) 기준 · 박람회 현장 계약 시 · 평형·현장 여건에 따라 달라질 수 있음"
m["A2"].font = f_note
header(m, 4, ["품목", "시공 범위", "일반", "일반 자재·내용", "프리미엄", "프리미엄 자재·내용", "비고·옵션"])
r = 5
master = {}  # item -> 첫 행(업체별 시트 기본 패키지)
cat2item = {}
for grp in data["groups"]:
    m.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
    cs(m, r, 1, grp["name"], f_grp, left, fill_grp)
    for c in range(2, 8):
        m.cell(r, c).border = box
        m.cell(r, c).fill = fill_grp
    r += 1
    for it in grp["items"]:
        master.setdefault(it["item"], r)
        for c in it["cats"]:
            cat2item.setdefault(c, it["item"])
        it["_row"] = r
        cs(m, r, 1, it["item"], f_item)
        cs(m, r, 2, it["scope"], f_input, left)
        for col, key in ((3, "normal"), (5, "premium")):
            v = it[key]
            cs(m, r, col, v, f_price, center, fill_yel if v is None and it["item"] == "조명" else None, PRICE)
        cs(m, r, 4, it["normal_desc"], f_input, left)
        cs(m, r, 6, it["premium_desc"], f_input, left)
        cs(m, r, 7, it["option"], f_input, left)
        m.row_dimensions[r].height = 34
        r += 1
M_END = r - 1
m.cell(r + 1, 1, "※ 기본 범위만 업체 공통. 추가 범위·옵션과 빈 프리미엄 칸은 업체가 박람회 현장에서 제시. 줄눈 일반 = 폴리우레아, 프리미엄 = 케라폭시.").font = f_note
m.cell(r + 2, 1, "※ 가전·가구는 업체별 대표 특가 1~2개 사전 공개, 그 외 전 품목 박람회 현장 공개 ('가전가구_특가' 시트)").font = f_note
m.cell(r + 3, 1, "※ 조명(노란 칸) = 엣지컴퍼니 직영 — 본부장님 입력").font = f_note
widths(m, [13, 38, 10, 14, 10, 24, 24])
m.freeze_panes = "A5"
a4(m)
m.print_title_rows = "4:4"

# ======================= 업체별_작성용 =======================
lineup = {}
for x in rows:
    if x["event"] == EVENT and x["status"] != "불참":
        lineup.setdefault(x["cat"], []).append(x)
w = wb.create_sheet("업체별_작성용")
w["A1"] = "업체별 시공 단가 (업체명 포함 · 내부 전용)"
w["A1"].font = f_title
w["A2"] = "명단 = 최근 행사(에코델타 디에트르) 참가 업체. 초록 = 시공단가표 대표 패키지(품목 첫 줄) 자동. 업체 가격이 다르면 그 칸에 덮어쓰기. 행 순서 = 공개 코드 A·B·C 순서."
w["A2"].font = f_note
header(w, 4, ["품목", "공개코드", "업체명", "시공 범위", "일반", "일반 자재", "프리미엄", "프리미엄 자재", "비고", "업체 확인"])
r = 5
W0 = r
for grp in data["groups"]:
    for it in grp["items"]:
        targets = []
        for cat in it["cats"]:
            if cat2item.get(cat) != it["item"]:
                continue
            targets += [(cat, v) for v in sorted(lineup.get(cat, []), key=lambda x: x["code"])
                        if v["vendor"] not in it.get("exclude_vendors", [])]
        targets += [(None, dict(vendor=x["vendor"], status="유료")) for x in it.get("extra_vendors", [])]
        if True:
            for cat, v in targets:
                mr = master[it["item"]]
                link = lambda col: f'=IF(시공단가표!{col}{mr}="","",시공단가표!{col}{mr})'
                cs(w, r, 1, it["item"], f_item)
                sep = " " if it["item"][-1].isascii() else ""
                cs(w, r, 2, f'=A{r}&"{sep}"&CHAR(64+COUNTIF($A${W0}:A{r},A{r}))', f_bold)
                name = "엣지컴퍼니(직영)" if cat == "조명" else v["vendor"]
                cs(w, r, 3, name, f_input, left)
                o = OVR.get(f'{it["item"]}|{name}', {})
                for col, key, src, lf, al in ((4, "scope", "B", f_link, left), (5, "normal", "C", f_linkb, center),
                                              (6, "ndesc", "D", f_link, left), (7, "premium", "E", f_linkb, center),
                                              (8, "pdesc", "F", f_link, left)):
                    if key in o:  # 본부장님 입력값 = 파란 글씨(직접 입력)
                        cs(w, r, col, o[key], f_price if col in (5, 7) else f_input, al, fmt=PRICE if col in (5, 7) else None)
                    else:
                        cs(w, r, col, link(src), lf, al, fmt=PRICE if col in (5, 7) else None)
                note = o.get("note") or ("직전 행사 무상 참가 — 참가 여부 확인" if v["status"] != "유료" else None)
                cs(w, r, 9, note, f_input, left)
                ok = bool(o) or it.get("confirmed")
                cs(w, r, 10, "본부장님 입력" if ok else "미확인", f_input, center, None if ok else fill_yel)
                w.row_dimensions[r].height = 30
                r += 1
W1 = r - 1
widths(w, [13, 14, 22, 34, 10, 12, 10, 18, 24, 10])
w.freeze_panes = "D5"
w.auto_filter.ref = f"A4:J{W1}"

# ======================= 업체별_공개용 =======================
p = wb.create_sheet("업체별_공개용")
p["A1"] = "입주박람회 시공 단가표 (업체별)"
p["A1"].font = f_title
p["A2"] = "84㎡(34평) 기준 · 박람회 현장 계약 시 · 업체명은 박람회 현장에서 공개"
p["A2"].font = f_note
header(p, 4, ["품목", "업체", "시공 범위", "일반", "일반 자재", "프리미엄", "프리미엄 자재", "비고"])
for i, rr in enumerate(range(W0, W1 + 1)):
    r = 5 + i
    ref = lambda col: f"업체별_작성용!{col}{rr}"
    bl = lambda col: f'=IF({ref(col)}="","",{ref(col)})'
    cs(p, r, 1, f"={ref('A')}", f_item)
    cs(p, r, 2, f"={ref('B')}", f_bold)
    cs(p, r, 3, bl("D"), f_base, left)
    cs(p, r, 4, f'=IF({ref("E")}="","현장 제시",{ref("E")})', f_item, fmt=PRICE)
    cs(p, r, 5, bl("F"), f_base, left)
    cs(p, r, 6, f'=IF({ref("G")}="","현장 제시",{ref("G")})', f_item, fmt=PRICE)
    cs(p, r, 7, bl("H"), f_base, left)
    cs(p, r, 8, f'=IF(LEFT({ref("I")},5)="직전 행사","",IF({ref("I")}="","",{ref("I")}))', f_base, left)
    p.row_dimensions[r].height = 30
widths(p, [13, 14, 34, 10, 12, 10, 18, 20])
p.freeze_panes = "C5"
a4(p)
p.print_title_rows = "4:4"

# ======================= 가전가구_특가 =======================
s = wb.create_sheet("가전가구_특가")
s["A1"] = "가전·가구 특가 — 업체별 1~2개 사전 공개, 나머지 박람회 현장 공개"
s["A1"].font = f_title
s["A2"] = "업체에서 받은 특가만 입력. 정상가는 실제 판매가여야 함(허위 할인율 금지). 한정수량은 숫자로. 공개용은 아래 '가전가구_공개용'."
s["A2"].font = f_note
header(s, 4, ["품목", "공개코드", "업체명", "특가① 상품", "정상가", "박람회 특가", "할인율", "한정", "특가② 상품", "정상가", "박람회 특가", "할인율", "한정"])
r = 5
S0 = r
for cat in data["special"]:
    for v in sorted(lineup.get(cat, []), key=lambda x: x["code"]):
        cs(s, r, 1, cat, f_item)
        cs(s, r, 2, f'=A{r}&CHAR(64+COUNTIF($A${S0}:A{r},A{r}))', f_bold)
        cs(s, r, 3, v["vendor"], f_input, left)
        for c in (4, 9):
            cs(s, r, c, None, f_input, left, fill_yel)
        for c in (5, 6, 10, 11):
            cs(s, r, c, None, f_input, center, fill_yel, PRICE)
        for c in (8, 13):
            cs(s, r, c, None, f_input, center, fill_yel)
        cs(s, r, 7, f'=IF(OR(E{r}="",F{r}="",E{r}=0),"",1-F{r}/E{r})', fmt=PCT)
        cs(s, r, 12, f'=IF(OR(J{r}="",K{r}="",J{r}=0),"",1-K{r}/J{r})', fmt=PCT)
        s.row_dimensions[r].height = 28
        r += 1
S1 = r - 1
widths(s, [10, 11, 20, 24, 10, 11, 7, 7, 24, 10, 11, 7, 7])
s.freeze_panes = "D5"

q = wb.create_sheet("가전가구_공개용")
q["A1"] = "가전·가구 박람회 특가"
q["A1"].font = f_title
q["A2"] = "업체별 대표 특가 사전 공개 · 그 외 전 품목 박람회 현장 공개 · 한정수량 소진 시 종료"
q["A2"].font = f_note
header(q, 4, ["품목", "업체", "특가① 상품", "정상가", "박람회 특가", "할인율", "특가② 상품", "정상가", "박람회 특가", "할인율", "그 외"])
for i, rr in enumerate(range(S0, S1 + 1)):
    r = 5 + i
    ref = lambda col: f"가전가구_특가!{col}{rr}"
    bl = lambda col: f'=IF({ref(col)}="","",{ref(col)})'
    cs(q, r, 1, f"={ref('A')}", f_item)
    cs(q, r, 2, f"={ref('B')}", f_bold)
    cs(q, r, 3, f'=IF({ref("D")}="","현장 공개",{ref("D")})', f_base, left)
    cs(q, r, 4, bl("E"), fmt=PRICE)
    cs(q, r, 5, bl("F"), f_item, fmt=PRICE)
    cs(q, r, 6, bl("G"), fmt=PCT)
    cs(q, r, 7, bl("I"), f_base, left)
    cs(q, r, 8, bl("J"), fmt=PRICE)
    cs(q, r, 9, bl("K"), f_item, fmt=PRICE)
    cs(q, r, 10, bl("L"), fmt=PCT)
    cs(q, r, 11, "전 품목 박람회 현장 공개", f_base, left)
    q.row_dimensions[r].height = 28
widths(q, [10, 11, 24, 10, 11, 7, 24, 10, 11, 7, 20])
a4(q, landscape=True)

# ======================= 시세근거 =======================
t = wb.create_sheet("시세근거")
t["A1"] = "대표가 근거 — 공개 시세 범위와 출처 (2026-10 웹 조사, 84㎡ 기준, 단위 만원)"
t["A1"].font = f_title
t["A2"] = "제시 = 본부장님 제시가 / 확인 = 출처로 확인한 시세 범위 안의 대표가 / 추정 = 공개 시세 부족, AI 추정. 박람회 현장 계약가는 공개 자료가 거의 없음 → 업체 확인 필요."
t["A2"].font = f_note
header(t, 4, ["품목", "일반(대표가)", "프리미엄(대표가)", "공개 시세 범위", "근거", "주요 출처"])
r = 5
for grp in data["groups"]:
    for it in grp["items"]:
        mr = it["_row"]
        cs(t, r, 1, it["item"] + (f" ({it['scope'].split('(')[-1]}" if it["item"] == "줄눈" else ""), f_item)
        cs(t, r, 2, f'=IF(시공단가표!C{mr}="","",시공단가표!C{mr})', f_link, fmt=PRICE)
        cs(t, r, 3, f'=IF(시공단가표!E{mr}="","",시공단가표!E{mr})', f_link, fmt=PRICE)
        cs(t, r, 4, it["ref"], f_base, left)
        cs(t, r, 5, it["basis"], f_bold)
        cs(t, r, 6, it["src"], f_base, left)
        t.row_dimensions[r].height = 32
        r += 1
widths(t, [14, 12, 13, 36, 9, 52])
t.freeze_panes = "B5"

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = {"시공단가표": NAVY, "업체별_공개용": "2E7D32", "가전가구_공개용": "2E7D32"}.get(ws.title, "8FC0F0")
wb.active = 1
wb.save(OUT)
print("saved", OUT, "vendor rows", W1 - W0 + 1, "special rows", S1 - S0 + 1)
