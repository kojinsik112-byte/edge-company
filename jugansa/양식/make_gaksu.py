"""입주박람회 특약 이행 각서 — 편집용 .docx(바탕체, 관공서 각서 형식). 사용: python make_gaksu.py <출력폴더>"""
import os, sys
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
F = "바탕"
def run(p, t, size=11, bold=False):
    r = p.add_run(t); r.font.size = Pt(size); r.bold = bold; r.font.name = F; r._element.rPr.rFonts.set(qn("w:eastAsia"), F); return r
def P(d, t="", size=11, bold=False, align=None, before=0, after=4):
    p = d.add_paragraph(); run(p, t, size, bold)
    if align is not None: p.alignment = align
    p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after); return p
def cell(c, t, bold=False, center=False, w=None):
    c.text = ""; p = c.paragraphs[0]; run(p, t, 10.5, bold)
    if center: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if w: c.width = Cm(w)
def main(out):
    d = Document()
    for s in d.sections: s.top_margin = s.bottom_margin = Cm(1.8); s.left_margin = s.right_margin = Cm(2.0)
    st = d.styles["Normal"]; st.font.name = F; st.element.rPr.rFonts.set(qn("w:eastAsia"), F); st.font.size = Pt(11)
    P(d, "특 약 이 행 각 서", 22, True, WD_ALIGN_PARAGRAPH.CENTER, after=12)
    t = d.add_table(rows=4, cols=4); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    rows = [("상    호", "주식회사 엣지컴퍼니", "대 표 자", "대표이사 고진식"),
            ("사업자등록번호", "508-81-42798", "법인등록번호", "230111-0369271"),
            ("주    소", "울산광역시 울주군 청량읍 온산로 615-1, 2층 · 3층", "", ""),
            ("건    명", "(                    ) 입주박람회 및 공동구매 주관", "", "")]
    W = [3.2, 6.0, 2.8, 5.0]
    for i, r in enumerate(rows):
        for j, v in enumerate(r): cell(t.cell(i, j), v, bold=(j in (0, 2)), center=(j in (0, 2)), w=W[j])
    for i in (2, 3):
        a = t.cell(i, 1).merge(t.cell(i, 3)); cell(a, rows[i][1])
    P(d, "당사는 (                    ) 입주예정자협의회(이하 “협의회”)와 20    년    월    일 체결한 「공동구매 및 입주박람회 "
         "기획·진행 위탁 협약」(이하 “협약”)에 따라, 협약 체결 시 약정한 아래의 특약 사항을 성실히 이행할 것을 각서합니다.", before=10, after=6)
    P(d, "1. 특약 사항", 11, True, after=3)
    t = d.add_table(rows=9, cols=4); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, (h, w) in enumerate([("번호", 1.3), ("특 약 내 용", 8.2), ("이 행 시 기", 3.6), ("비 고", 3.9)]): cell(t.cell(0, j), h, True, True, w)
    for i in range(1, 9):
        cell(t.cell(i, 0), str(i), center=True, w=1.3)
        for j, w in ((1, 8.2), (2, 3.6), (3, 3.9)): cell(t.cell(i, j), "", w=w)
    P(d, "2. 이행 조건", 11, True, before=10, after=3)
    for i, s in enumerate([
        "위 특약 사항은 협약과 같은 효력을 가지며, 당사는 정해진 시기에 이를 이행하고 그 결과(사진·영수증 등)를 협의회에 서면으로 보고합니다.",
        "특약 사항의 내용이나 이행 시기를 변경할 때에는 사전에 협의회와 서면으로 합의합니다.",
        "특약 이행에 드는 비용은 당사가 전액 부담하며, 이를 공동구매 품목의 가격에 전가하지 않습니다.",
        "당사가 정당한 사유 없이 특약 사항을 이행하지 않는 경우, 협의회의 서면 시정 요구를 받은 날부터 (    )일 이내에 이행하며, 그 기간 안에 이행하지 않으면 협약 해지 사유가 됨을 인정합니다.",
        "특약 사항의 미이행으로 협의회 또는 입주예정자에게 손해가 발생한 경우 당사는 그 손해를 배상합니다.",
        "이 각서의 해석에 이견이 있는 경우 관계 법령과 협약에 반하지 않는 범위에서 협의회의 해석에 따릅니다."], 1):
        P(d, f"  {i}. {s}", 10.5, after=2)
    P(d, "위 사항을 성실히 이행할 것을 각서합니다.", 11, False, WD_ALIGN_PARAGRAPH.CENTER, before=12, after=6)
    P(d, "20    년    월    일", 11, False, WD_ALIGN_PARAGRAPH.CENTER, after=8)
    for s in ["각  서  인", "상      호 :  주식회사 엣지컴퍼니", "대  표  자 :  대표이사  고 진 식            (인)"]:
        p = P(d, s, 11, s.startswith("각"), after=2); p.paragraph_format.left_indent = Cm(7.5)
    P(d, "붙임  1. 공동구매 및 입주박람회 기획·진행 위탁 협약서 사본 1부.", 10, before=10, after=0)
    P(d, "        2. 법인인감증명서 1부.  끝.", 10, after=8)
    P(d, "(                    ) 입주예정자협의회 회장 귀하", 13, True, WD_ALIGN_PARAGRAPH.CENTER, before=6)
    d.save(os.path.join(out, "엣지컴퍼니_입주박람회_특약이행각서_양식.docx"))
if __name__ == "__main__":
    o = sys.argv[1] if len(sys.argv) > 1 else "."; os.makedirs(o, exist_ok=True); main(o); print("ok")
