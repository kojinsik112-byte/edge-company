"""입주박람회 주관 협약서(표준) + 주관 실적 확인서 — 편집 가능한 .docx 양식.
사용: python make_hyupyak.py <출력폴더>
- 협약서: 앞으로 주관 단지와 체결할 표준 협약서(빈칸 채워 2부 작성·서명).
- 실적 확인서: 협약서를 보관하지 않은 과거 단지용. 당시 협의회 대표가 '지금 날짜'로 서명하는 사실 확인서.
  과거 날짜로 협약서를 새로 만들어 제출하지 않는다(사문서 소급 작성 = 입찰 무효·형사 위험)."""
import os, sys
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

CO = dict(name="주식회사 엣지컴퍼니", ceo="고진식", brn="508-81-42798",
          addr="울산광역시 울주군 청량읍 온산로 615-1, 2·3층", tel="010-7273-9901 (주관사업 총괄 본부장)")
FONT = "맑은 고딕"
NAVY = RGBColor(0x0D, 0x1E, 0x33)

def base():
    d = Document()
    for s in d.sections:
        s.top_margin = s.bottom_margin = Cm(2.0); s.left_margin = s.right_margin = Cm(2.2)
    st = d.styles["Normal"]; st.font.name = FONT; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.paragraph_format.space_after = Pt(3); st.paragraph_format.line_spacing = 1.25
    return d

def para(d, text="", size=None, bold=False, align=None, color=None, after=None, before=None):
    p = d.add_paragraph(); r = p.add_run(text); r.bold = bold
    if size: r.font.size = Pt(size)
    if color: r.font.color.rgb = color
    if align: p.alignment = align
    if after is not None: p.paragraph_format.space_after = Pt(after)
    if before is not None: p.paragraph_format.space_before = Pt(before)
    return p

def title(d, t, sub=None):
    para(d, t, 20, True, WD_ALIGN_PARAGRAPH.CENTER, NAVY, after=2)
    if sub: para(d, sub, 10, False, WD_ALIGN_PARAGRAPH.CENTER, RGBColor(0x55, 0x5B, 0x66), after=10)

def table(d, rows, widths, head=False):
    t = d.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = t.cell(i, j); c.width = Cm(widths[j]); c.text = ""
            r = c.paragraphs[0].add_run(v); r.font.size = Pt(10)
            if (head and i == 0) or (not head and j % 2 == 0 and len(row) > 1 and v):
                r.bold = True
    d.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def article(d, no, head, lines):
    para(d, f"제{no}조 ({head})", 11, True, before=6, after=2)
    for ln in lines: para(d, ln, after=1)

def sign_block(d, left_title, right_title):
    t = d.add_table(rows=6, cols=2); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    L = [left_title, "단 체 명 :", "대 표 자 :                      (서명 또는 인)", "직    위 :", "연 락 처 :", "주    소 :"]
    R = [right_title, f"상    호 : {CO['name']}", f"대 표 자 : 대표이사 {CO['ceo']}        (인)",
         f"사업자등록번호 : {CO['brn']}", f"연 락 처 : {CO['tel']}", f"주    소 : {CO['addr']}"]
    for i in range(6):
        for j, v in enumerate((L[i], R[i])):
            c = t.cell(i, j); c.width = Cm(8.3); c.text = ""
            r = c.paragraphs[0].add_run(v); r.font.size = Pt(9.5); r.bold = (i == 0)

def hyupyak(out):
    d = base()
    title(d, "입주박람회 주관사 협약서", "(공동구매 · 입주박람회 위탁 운영)")
    para(d, "○○○○ 입주예정자협의회(이하 \"협의회\")와 주식회사 엣지컴퍼니(이하 \"주관사\")는 아래 단지의 입주박람회 및 "
            "공동구매 주관 업무에 관하여 다음과 같이 협약을 체결한다.", after=6)
    table(d, [["단 지 명", "", "소 재 지", ""], ["세 대 수", "            세대", "입주 예정", "      년      월"],
              ["시공(시행)사", "", "선정 방식", "□ 공개 입찰  □ 기타(          )"]], [2.6, 5.7, 2.6, 5.7])
    article(d, 1, "목적", ["이 협약은 협의회가 주관사에 입주박람회와 공동구매 운영을 맡기고, 입주예정자에게 투명한 가격과 "
                          "안정적인 시공·A/S를 제공하기 위하여 양 당사자의 권리와 의무를 정함을 목적으로 한다."])
    article(d, 2, "협약 기간", ["협약 체결일부터 공동구매 품목의 무상 A/S 기간이 끝나는 날까지로 한다. "
                              "단, 행사 운영 업무는 입주 지정기간 종료일까지로 하고, 이후에는 A/S 창구 업무만 수행한다."])
    article(d, 3, "주관사의 업무", [
        "① 입주박람회 기획·운영(행사장 대관, 부스 설치, 안내·주차·안전 인력 배치, 홍보물 제작).",
        "② 품목별 참여 업체 모집 및 심사 자료 준비 — 업체 선정은 협의회와 함께 공개 심사로 한다.",
        "③ 사전 공개 단가 확정 및 시중가·온라인 최저가 대비 가격 검수.",
        "④ 계약·시공·A/S 통합 창구 운영 및 입주민 민원 처리(접수 후 처리 결과 회신).",
        "⑤ 행사 결과 보고(방문 세대, 품목별 계약 건수, 민원 처리 현황)를 행사 종료 후 30일 이내 협의회에 제출."])
    article(d, 4, "협의회의 업무", ["① 입주예정자 대상 행사 공지 및 일정 협의. ② 업체 선정 심사 참여. "
                                  "③ 주관사 업무 수행에 필요한 범위의 협조. 협의회는 주관사의 업무 수행에 대한 비용을 부담하지 않는다."])
    article(d, 5, "행사 일정 및 장소", ["행사 기간 :      년    월    일 ~    월    일 (   일간)    장소 :",
                                    "일정과 장소는 양 당사자가 협의하여 확정하며, 변경 시 서면(전자우편 포함)으로 합의한다."])
    article(d, 6, "입주민 혜택", ["입주민 혜택의 종류·규모·지급 방법은 별도 합의서(부속서)로 정하며, 공동구매 단가에 전가하지 않는다."])
    article(d, 7, "계약 취소·환불 및 A/S", ["① 품목별 취소·환불 기준은 주관사가 제시하고 협의회가 승인한 '계약 취소·환불 규정'에 따른다.",
                                        "② 시공 하자 발생 시 주관사는 해당 업체에 보수를 요청하고, 업체가 이행하지 않으면 주관사가 책임지고 처리한다."])
    article(d, 8, "안전 및 보험", ["주관사는 행사 기간·장소에 맞춰 행사주최자배상책임보험에 가입하고, 협의회의 요청이 있으면 가입증명서를 제출한다."])
    article(d, 9, "금지 행위", ["① 주관사와 참여 업체는 협의회 임원·운영위원 개인에게 금품·향응을 제공하지 않는다.",
                              "② 협의회 임원은 업체 선정과 관련하여 개인적 이익을 요구하거나 받지 않는다."])
    article(d, 10, "개인정보 보호 및 비밀 유지", ["입주예정자의 개인정보는 행사 운영 목적으로만 이용하고, 목적 달성 후 지체 없이 파기한다. "
                                              "양 당사자는 협약 과정에서 알게 된 상대방의 비밀을 제3자에게 누설하지 않는다."])
    article(d, 11, "협약 해지", ["어느 한쪽이 이 협약을 위반하고 서면 시정 요구 후 14일 이내에 시정하지 않으면 상대방은 서면으로 협약을 해지할 수 있다. "
                               "해지 전에 체결된 입주민 개별 계약과 A/S 의무는 해지 후에도 유지된다."])
    article(d, 12, "분쟁 해결", ["이 협약에 관한 분쟁은 상호 협의로 해결하고, 협의가 이루어지지 않으면 협의회 소재지 관할 법원을 관할 법원으로 한다."])
    article(d, 13, "기타", ["이 협약에 정하지 않은 사항은 양 당사자가 협의하여 서면으로 정한다. 이 협약서는 2부를 작성하여 각 1부씩 보관한다.",
                          "※ 협약서와 부속서는 주관 실적 증빙으로 쓰이므로 주관사는 원본을 영구 보관한다."])
    para(d, "20    년      월      일", 11, True, WD_ALIGN_PARAGRAPH.CENTER, before=8, after=8)
    sign_block(d, "[ 협 의 회 ]", "[ 주 관 사 ]")
    d.save(os.path.join(out, "엣지컴퍼니_입주박람회_주관사_협약서_표준양식.docx"))

def hwagin(out):
    d = base()
    title(d, "입주박람회 주관 실적 확인서", "(협약서 원본 미보관 단지 — 당시 협의회 대표 사실 확인)")
    para(d, "■ 주관사", 11, True, after=2)
    table(d, [["상    호", CO["name"], "대 표 자", f"대표이사 {CO['ceo']}"],
              ["사업자등록번호", CO["brn"], "소 재 지", CO["addr"]]], [2.8, 5.5, 2.3, 6.0])
    para(d, "■ 행사 내용", 11, True, after=2)
    table(d, [["단 지 명", "", "소 재 지", ""], ["세 대 수", "            세대", "입주 연월", "      년      월"],
              ["행사 기간", "      년    월    일 ~    월    일", "행사 장소", ""],
              ["선정 방식", "□ 공개 입찰  □ 협의회 지정  □ 기타", "협약 체결", "      년    월경  □ 서면  □ 기타"],
              ["방문 세대(선택)", "약          세대", "참여 업체(선택)", "약          개 업체"]], [2.8, 5.5, 2.3, 6.0])
    para(d, "위 단지 입주예정자협의회는 위 기간에 주식회사 엣지컴퍼니를 입주박람회 주관사로 선정·협약하여 "
            "입주박람회와 공동구매를 진행하였음을 확인합니다.", before=4, after=2)
    para(d, "※ 협약서 원본은 행사 종료 후 보관하지 않아 이 확인서로 협약 사실을 확인합니다. 이 확인서는 아래 작성일에 작성되었습니다.",
         9.5, color=RGBColor(0x55, 0x5B, 0x66), after=2)
    para(d, "※ 확인자는 이 확인서와 연락처가 입찰 심사 기관(입주예정자협의회)에 실적 증빙으로 제출되고, "
            "사실 확인을 위한 연락에 쓰이는 것에 동의합니다.   □ 동의함", 9.5, color=RGBColor(0x55, 0x5B, 0x66), after=8)
    para(d, "작성일 :  20    년      월      일", 11, True, WD_ALIGN_PARAGRAPH.CENTER, before=6, after=8)
    t = d.add_table(rows=5, cols=2); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    rows = [("[ 확 인 자 ]", ""), ("단 체 명 (당시)", "                    입주예정자협의회"),
            ("직위 (당시) / 성명", "                        /                    (서명 또는 인)"),
            ("연 락 처", ""), ("현재 직위 (해당 시)", "□ 입주자대표회의 (        )   □ 해당 없음")]
    for i, (a, b) in enumerate(rows):
        for j, v in enumerate((a, b)):
            c = t.cell(i, j); c.width = Cm([4.5, 12.1][j]); c.text = ""
            r = c.paragraphs[0].add_run(v); r.font.size = Pt(10); r.bold = (j == 0)
    d.save(os.path.join(out, "엣지컴퍼니_입주박람회_주관실적_확인서_양식.docx"))

if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "."
    os.makedirs(out, exist_ok=True); hyupyak(out); hwagin(out); print("ok", out)
