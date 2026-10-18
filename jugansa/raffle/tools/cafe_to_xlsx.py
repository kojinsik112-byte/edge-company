"""네이버 카페 댓글 → 추첨기용 엑셀 변환기 (엣지컴퍼니 주관사)

사용법
  python cafe_to_xlsx.py 입력.txt 출력.xlsx

입력 형식 (둘 다 자동 인식)
  1) 카페 댓글 화면을 그대로 복사한 텍스트
     (닉네임 줄 → 댓글 줄 → "2026.09.18. 16:21답글쓰기" 줄 반복, 마크다운 링크 섞여도 OK)
  2) 정리 형식 — 화답 댓글까지 넣을 때
     @닉네임|2026.09.18. 16:21
     댓글 내용(여러 줄 가능)
     >화답 댓글

출력 시트
  참여자        : 동호수 | 닉네임 | 댓글 내용 | 화답 댓글   ← 추첨기에 이 파일 그대로 넣기
  원본·작성일시 : 작성자 원문, 작성일시까지
  검토 필요     : 같은 닉네임 중복 · 한 세대 여러 계정 · 동호수 없음 · 주관사 댓글
"""
import re
import sys
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

DATE_RE = re.compile(r'^(\d{4}\.\s?\d{1,2}\.\s?\d{1,2}\.?\s+\d{1,2}:\d{2})(.*)$')
JUNK_RE = re.compile(r'^(\d{1,4}|답글쓰기|등록순|최신순|새로고침|댓글|.*클린봇.*)$')
STAFF_RE = re.compile(r'엣지\s*컴퍼니')


def split_unit(author):
    """'110동 1201호 하늘맘B' / '1051404' / '105 2002 별빛A' → ('110동 1201호', '하늘맘B')"""
    s = re.sub(r'\s+', ' ', author).strip()
    pats = [
        r'^(\d{3})\s*동\s*(\d{2,4})\s*호?\s*(.*)$',   # 110동 1201호 닉 / 109동903닉
        r'^(\d{3})\s+(\d{3,4})(?!\d)\s*(.*)$',        # 105 2002 닉
        r'^(\d{3})(\d{3,4})(?!\d)\s*(.*)$',           # 1051404닉 / 101101 닉
    ]
    for p in pats:
        m = re.match(p, s)
        if m:
            ho = int(m.group(2))
            if 101 <= ho <= 4999 and 1 <= ho % 100 <= 12:
                return f'{int(m.group(1))}동 {ho}호', m.group(3).strip()
    return '', s


def parse_blocks(text):
    """정리 형식(@작성자|일시 / 댓글 / >화답)"""
    out, cur = [], None
    for line in text.splitlines():
        if line.startswith('@'):
            if cur:
                out.append(cur)
            author, _, when = line[1:].partition('|')
            cur = {'author': author.strip(), 'when': when.strip(), 'comment': [], 'reply': ''}
        elif line.startswith('>') and cur:
            cur['reply'] = line[1:].strip()
        elif cur is not None:
            cur['comment'].append(line)
    if cur:
        out.append(cur)
    return out


def parse_cafe(text):
    """카페 화면 복사본(마크다운 링크 포함/미포함)"""
    text = re.sub(r'\[([^\]]*)\]\(https?://[^)]*\)', r'\1', text)
    lines = [re.sub(r'^\*\s+', '', l).rstrip() for l in text.splitlines()]
    out, author, buf = [], None, []
    for line in lines:
        m = DATE_RE.match(line.strip())
        if m:
            if author:
                out.append({'author': author, 'when': m.group(1), 'comment': buf, 'reply': ''})
            author, buf = None, []
            rest = m.group(2).replace('답글쓰기', '').strip()
            if rest and not JUNK_RE.match(rest):
                author = rest  # 복사 과정에서 다음 작성자가 날짜 줄에 붙은 경우
            continue
        if author is None:
            if line.strip() and not JUNK_RE.match(line.strip()):
                author = line.strip()
        else:
            buf.append(line)
    return out


def load(text):
    items = parse_blocks(text) if re.search(r'^@', text, re.M) else parse_cafe(text)
    rows, staff = [], []
    for it in items:
        comment = '\n'.join(it['comment']).strip()
        comment = re.sub(r'\n{3,}', '\n\n', comment)
        if STAFF_RE.search(it['author']):
            staff.append(it)
            continue
        unit, nick = split_unit(it['author'])
        rows.append({**it, 'unit': unit, 'nick': nick, 'comment': comment})
    return rows, staff


def style_header(ws, widths):
    head = Font(bold=True, color='FFFFFF')
    fill = PatternFill('solid', fgColor='10243F')
    for i, w in enumerate(widths, 1):
        c = ws.cell(row=1, column=i)
        c.font, c.fill, c.alignment = head, fill, Alignment(horizontal='center', vertical='center')
        ws.column_dimensions[c.column_letter].width = w
    ws.freeze_panes = 'A2'
    ws.row_dimensions[1].height = 22


def build(rows, staff, path):
    wb = Workbook()
    ws = wb.active
    ws.title = '참여자'
    ws.append(['동호수', '닉네임', '댓글 내용', '화답 댓글'])
    wrap = Alignment(wrap_text=True, vertical='top')
    for r in rows:
        ws.append([r['unit'], r['nick'], r['comment'], r['reply']])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = wrap
    style_header(ws, [15, 18, 70, 70])

    ws2 = wb.create_sheet('원본·작성일시')
    ws2.append(['번호', '작성자(원문)', '동호수', '닉네임', '작성일시', '댓글 내용'])
    for i, r in enumerate(rows, 1):
        ws2.append([i, r['author'], r['unit'], r['nick'], r['when'], r['comment']])
    for row in ws2.iter_rows(min_row=2):
        for c in row:
            c.alignment = wrap
    style_header(ws2, [7, 24, 15, 18, 18, 70])

    ws3 = wb.create_sheet('검토 필요')
    ws3.append(['구분', '동호수', '닉네임', '작성일시', '비고'])
    by_author, by_unit = defaultdict(list), defaultdict(list)
    for r in rows:
        by_author[(r['unit'], r['nick'])].append(r)
        if r['unit']:
            by_unit[r['unit']].append(r)
    issues = 0
    for (unit, nick), rs in by_author.items():
        if len(rs) > 1:
            issues += 1
            ws3.append(['같은 닉네임 중복 댓글', unit, nick, ', '.join(x['when'] for x in rs), f'{len(rs)}회 작성 → 추첨기 "1인 1회"에서 자동으로 1명 처리'])
    for unit, rs in by_unit.items():
        nicks = sorted({x['nick'] for x in rs})
        if len(nicks) > 1:
            issues += 1
            ws3.append(['한 세대 여러 계정', unit, ' / '.join(nicks), '', '이벤트 규칙이 "세대당 1회"면 추첨기에서 "세대당 1회" 선택'])
    for r in rows:
        if not r['unit']:
            issues += 1
            ws3.append(['동호수 없음', '', r['nick'], r['when'], '닉네임에서 동호수를 찾지 못함 — 확인 필요'])
        elif not r['nick']:
            ws3.append(['닉네임 없음', r['unit'], '', r['when'], '추첨기에는 "입주민"으로 표시'])
    for s in staff:
        ws3.append(['주관사 댓글(제외)', '', s['author'], s['when'], '추첨 대상에서 제외함'])
    style_header(ws3, [22, 15, 30, 40, 52])
    wb.save(path)
    return by_author, by_unit, issues


def main():
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(1)
    src, dst = sys.argv[1], sys.argv[2]
    rows, staff = load(open(src, encoding='utf-8').read())
    by_author, by_unit, _ = build(rows, staff, dst)
    print(f'댓글 {len(rows)}건 (주관사 {len(staff)}건 제외) → {dst}')
    print(f'  1인 1회 기준 추첨 대상: {len(by_author)}명')
    print(f'  세대당 1회 기준 추첨 대상: {len(by_unit) + sum(1 for r in rows if not r["unit"])}세대')
    print(f'  동호수 인식 실패: {sum(1 for r in rows if not r["unit"])}건 · 닉네임 없음: {sum(1 for r in rows if r["unit"] and not r["nick"])}건')


if __name__ == '__main__':
    main()
