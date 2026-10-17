# -*- coding: utf-8 -*-
"""docx → HTML 미리보기(LibreOffice가 이 컨테이너에서 안 돌아 Chromium으로 쪽 모양을 확인하기 위한 보조 · 글꼴은 Pretendard 대체)."""
import html, sys
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

def runs(p):
    out = ""
    for r in p.runs:
        rpr = r._element.rPr
        st = []
        if r.font.size: st.append(f"font-size:{r.font.size.pt}pt")
        if r.font.bold: st.append("font-weight:700")
        if r.font.color is not None and r.font.color.rgb is not None: st.append(f"color:#{r.font.color.rgb}")
        if rpr is not None:
            shd = rpr.find(qn("w:shd"))
            if shd is not None: st.append(f"background:#{shd.get(qn('w:fill'))};padding:0 2px")
            if rpr.find(qn("w:br")) is not None: out += "<br>"
        for br in r._element.findall(qn("w:br")):
            if br.get(qn("w:type")) == "page": out += '<div class="pb"></div>'
        out += f'<span style="{";".join(st)}">{html.escape(r.text)}</span>'
    return out

def para(p):
    pf = p.paragraph_format; st = []
    if pf.space_before is not None: st.append(f"margin-top:{pf.space_before.pt}pt")
    if pf.space_after is not None: st.append(f"margin-bottom:{pf.space_after.pt}pt")
    if pf.line_spacing is not None: st.append(f"line-height:{pf.line_spacing}")
    if pf.left_indent is not None: st.append(f"margin-left:{pf.left_indent.cm}cm")
    if pf.first_line_indent is not None: st.append(f"text-indent:{pf.first_line_indent.cm}cm")
    if p.alignment is not None: st.append("text-align:" + {1: "center", 2: "right", 3: "justify"}.get(int(p.alignment), "left"))
    ppr = p._element.pPr
    if ppr is not None and ppr.find(qn("w:pBdr")) is not None:
        b = ppr.find(qn("w:pBdr")).find(qn("w:bottom"))
        st.append(f"border-bottom:{int(b.get(qn('w:sz')))/8}pt solid #{b.get(qn('w:color'))}")
    return f'<p style="{";".join(st)}">{runs(p) or "&nbsp;"}</p>'

def table(t):
    out = '<table>'
    for row in t.rows:
        out += "<tr>"
        seen = set()
        for c in row.cells:
            if id(c._element) in seen: continue
            seen.add(id(c._element))
            tcpr = c._element.tcPr; st = []
            if tcpr is not None:
                shd = tcpr.find(qn("w:shd"))
                if shd is not None: st.append(f"background:#{shd.get(qn('w:fill'))}")
                span = tcpr.find(qn("w:gridSpan"))
            if c.width: st.append(f"width:{c.width.cm}cm")
            colspan = ""
            if tcpr is not None and tcpr.find(qn("w:gridSpan")) is not None: colspan = f' colspan="{tcpr.find(qn("w:gridSpan")).get(qn("w:val"))}"'
            out += f'<td{colspan} style="{";".join(st)}">' + "".join(para(p) for p in c.paragraphs) + "</td>"
        out += "</tr>"
    return out + "</table>"

doc = Document(sys.argv[1])
body = ""
for el in doc.element.body.iterchildren():
    if el.tag == qn("w:p"): body += para(Paragraph(el, doc))
    elif el.tag == qn("w:tbl"): body += table(Table(el, doc))
hdr = "".join(para(p) for p in doc.sections[0].header.paragraphs)
css = """@page{size:A4;margin:2.2cm 2cm 2cm 2cm}body{font-family:Pretendard,'Malgun Gothic',sans-serif;font-size:10.5pt;color:#111;margin:0}
p{margin:0;word-break:keep-all}table{border-collapse:collapse;width:100%;margin:0 auto;table-layout:fixed}td{border:.6pt solid #C9CED6;padding:3pt 5.5pt;vertical-align:top}
.pb{page-break-before:always}.hdr{position:running(h);font-size:8.5pt;color:#5B6676;text-align:right}"""
open(sys.argv[2], "w", encoding="utf-8").write(f"<!doctype html><meta charset=utf-8><style>{css}</style><div class=hdr>{hdr}</div>{body}")
print("ok")
